"""
Report generation service — PDF (reportlab) and Excel (openpyxl)
"""

import io
from datetime import datetime


def generate_pdf(event, plan) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=20, spaceAfter=6)
    h2_style = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, spaceAfter=4, textColor=colors.HexColor("#1e3a5f"))
    body_style = styles["BodyText"]

    plan_data = plan.plan_data or {}
    allocation = plan_data.get("resource_allocation", [])
    alerts = plan_data.get("alerts", [])
    recommendations = plan_data.get("recommendations", [])
    venues = plan_data.get("venues", [])
    transport = plan_data.get("transport_schedule", [])
    timeline = plan_data.get("timeline", [])

    story = []

    # Header
    story.append(Paragraph("Event Logistics Plan", title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}", body_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1e3a5f")))
    story.append(Spacer(1, 0.4*cm))

    # Event Details
    story.append(Paragraph("Event Details", h2_style))
    details_data = [
        ["Event Name", event.name],
        ["Type", event.event_type.value.title()],
        ["Date", event.date.strftime("%d %B %Y")],
        ["Duration", f"{event.duration_hours} hours"],
        ["Registrations", str(event.registrations)],
        ["Teams", str(event.teams or "N/A")],
    ]
    details_table = Table(details_data, colWidths=[5*cm, 11*cm])
    details_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.HexColor("#f0f4f8"), colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1dce8")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(details_table)
    story.append(Spacer(1, 0.5*cm))

    # Attendance Prediction
    story.append(Paragraph("Attendance Prediction", h2_style))
    story.append(Paragraph(
        f"<b>Predicted Attendance:</b> {plan.predicted_attendance} participants",
        body_style
    ))
    story.append(Paragraph(
        f"<b>Expected Range:</b> {plan.confidence_low} – {plan.confidence_high}",
        body_style
    ))
    story.append(Spacer(1, 0.5*cm))

    # Venue Allocation
    if venues:
        story.append(Paragraph("Venue Allocation", h2_style))
        venue_data = [["Venue", "Capacity", "Type"]] + [
            [v.get("name", ""), str(v.get("capacity", "")), v.get("venue_type", "")]
            for v in venues
        ]
        venue_table = Table(venue_data, colWidths=[8*cm, 4*cm, 4*cm])
        venue_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1dce8")),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(venue_table)
        story.append(Spacer(1, 0.5*cm))

    # Resource Allocation
    story.append(Paragraph("Resource Allocation", h2_style))
    res_data = [["Resource", "Required", "Available", "Allocated", "Shortage"]] + [
        [
            r.get("resource", "").title(),
            str(r.get("required", 0)),
            str(r.get("available", 0)),
            str(r.get("allocated", 0)),
            str(r.get("shortage", 0)) if r.get("shortage", 0) > 0 else "—",
        ]
        for r in allocation
    ]
    res_table = Table(res_data, colWidths=[4.5*cm, 3.5*cm, 3.5*cm, 3.5*cm, 2*cm])
    res_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1dce8")),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("TEXTCOLOR", (4, 1), (4, -1), colors.HexColor("#c0392b")),
    ]))
    story.append(res_table)
    story.append(Spacer(1, 0.5*cm))

    # Transport Schedule
    if transport:
        story.append(Paragraph("Transport Schedule", h2_style))
        for t in transport:
            story.append(Paragraph(f"• {t['time']} — {t['buses']}: {t['route']}", body_style))
        story.append(Spacer(1, 0.4*cm))

    # Timeline
    if timeline:
        story.append(Paragraph("Event Timeline", h2_style))
        for t in timeline:
            story.append(Paragraph(f"• {t['time']} — {t['activity']}", body_style))
        story.append(Spacer(1, 0.4*cm))

    # Alerts & Recommendations
    if alerts or recommendations:
        story.append(Paragraph("Alerts & Recommendations", h2_style))
        for a in alerts:
            icon = "⚠" if a["level"] == "warning" else ("✗" if a["level"] == "error" else "✓")
            story.append(Paragraph(f"{icon} {a['message']}: {a['detail']}", body_style))
        for rec in recommendations:
            story.append(Paragraph(f"→ {rec}", body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def generate_excel(event, plan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    plan_data = plan.plan_data or {}

    # ── Sheet 1: Summary ──
    ws = wb.active
    ws.title = "Summary"
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    header_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30

    ws.append(["Event Logistics Plan"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated", datetime.now().strftime("%d %B %Y, %H:%M")])
    ws.append([])
    ws.append(["EVENT DETAILS"])
    ws["A4"].font = bold
    for row in [
        ("Event Name", event.name),
        ("Type", event.event_type.value.title()),
        ("Date", event.date.strftime("%d %B %Y")),
        ("Duration", f"{event.duration_hours} hours"),
        ("Registrations", event.registrations),
        ("Teams", event.teams or "N/A"),
    ]:
        ws.append(row)
    ws.append([])
    ws.append(["ATTENDANCE PREDICTION"])
    ws[f"A{ws.max_row}"].font = bold
    ws.append(["Predicted Attendance", plan.predicted_attendance])
    ws.append(["Confidence Low", plan.confidence_low])
    ws.append(["Confidence High", plan.confidence_high])

    # ── Sheet 2: Resources ──
    ws2 = wb.create_sheet("Resource Allocation")
    ws2.column_dimensions["A"].width = 18
    for col in ["B", "C", "D", "E"]:
        ws2.column_dimensions[col].width = 14

    headers = ["Resource", "Required", "Available", "Allocated", "Shortage"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in plan_data.get("resource_allocation", []):
        ws2.append([
            r.get("resource", "").title(),
            r.get("required", 0),
            r.get("available", 0),
            r.get("allocated", 0),
            r.get("shortage", 0),
        ])

    # ── Sheet 3: Venues ──
    ws3 = wb.create_sheet("Venue Allocation")
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 15
    ws3.append(["Venue", "Capacity", "Type"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
    for v in plan_data.get("venues", []):
        ws3.append([v.get("name", ""), v.get("capacity", ""), v.get("venue_type", "")])

    # ── Sheet 4: Timeline ──
    ws4 = wb.create_sheet("Timeline")
    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 35
    ws4.append(["Time", "Activity"])
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font
    for t in plan_data.get("timeline", []):
        ws4.append([t.get("time", ""), t.get("activity", "")])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
